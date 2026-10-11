"""教案 → PPT（.pptx）。

PPT 是线性版面：没有纸面表格，只按模板的行结构把「标签 + 内容」排成文本框，
一屏放不下就整行续到下一页。没填的栏目直接跳过——纸面模板要留空白填写栏，幻灯片不需要。

文本框不支持行内样式，内容一律先洗成纯文本，否则 `<p>` 会当普通文字印出来。
"""

from __future__ import annotations

import io
import math
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.presentation import Presentation as PptxPresentation
from pptx.util import Inches, Pt

from ailan.ai_chat.files.layout import LayoutField, build_layout
from ailan.ai_chat.files.rich_text import field_value_to_text, to_plain_text_fields

SLIDE_W = 13.333
SLIDE_H = 7.5
MARGIN = 0.6
TEXT_W = SLIDE_W - MARGIN * 2
# 正文可用区：超过底部的行整条挪到下一页
BODY_TOP = 1.2
BODY_BOTTOM = SLIDE_H - 0.5
GUTTER = 0.25

TITLE_SIZE = 30
INFO_SIZE = 15
LABEL_SIZE = 15
BODY_SIZE = 13
# 一行文字的近似高度（英寸），估行高与判断翻页都用它
LINE_H = 0.3

DARK = RGBColor(0x33, 0x33, 0x33)
MUTED = RGBColor(0x66, 0x66, 0x66)
BLACK = RGBColor(0, 0, 0)


def _add_text(
    slide: Any,
    text: str,
    *,
    x: float,
    y: float,
    w: float,
    h: float,
    size: int,
    bold: bool = False,
    color: RGBColor = DARK,
    align_center: bool = False,
) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    if align_center:
        paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def _estimated_lines(text: str, width_in: float, size_pt: int) -> int:
    """按「汉字见方」估行数：PPT 的行高由渲染器决定，导出只能估——宁可估多不可估少。"""
    chars_per_line = max(4, int(width_in / (size_pt / 72)))
    return max(1, math.ceil(len(text) / chars_per_line))


def _cover(prs: PptxPresentation, title: str, info_line: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_text(slide, title, x=MARGIN, y=2.2, w=TEXT_W, h=1.2, size=TITLE_SIZE, bold=True, align_center=True)
    if info_line:
        _add_text(
            slide, info_line, x=MARGIN, y=3.5, w=TEXT_W, h=0.6, size=INFO_SIZE, color=MUTED, align_center=True
        )


def _filled_rows(fields: list[LayoutField], values: dict[str, Any], labels: dict[str, str]) -> list[Any]:
    """只留有内容的行：PPT 里没有「空白填写栏」这一说。"""
    layout = build_layout(fields, to_plain_text_fields(values), labels)
    filled = []
    for row in layout:
        kept = [item for item in row.fields if item.is_label or item.value.strip()]
        if kept:
            row.fields = kept
            filled.append(row)
    return filled


def _append_body(prs: PptxPresentation, rows: list[Any]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    y = BODY_TOP
    for row in rows:
        row_width = sum(item.width for item in row.fields)
        if row_width <= 0:
            continue
        usable = TEXT_W - GUTTER * (len(row.fields) - 1)

        # 先估每一列的高度，整行取最大值，否则多列会在纵向上叠在一起
        columns = []
        for item in row.fields:
            total = usable * (item.width / row_width)
            # 标签在左才并排占一列宽度；标签在上与无标签都独占整列
            label_w = min(1.6, total * 0.4) if item.label_position == "left" else 0.0
            body_w = total - label_w
            lines = 1 if item.is_label else _estimated_lines(item.value, body_w, BODY_SIZE)
            label_h = 0.32 if (not item.is_label and item.label_position == "top") else 0.0
            columns.append((item, total, label_w, body_w, label_h, lines))

        row_height = max(label_h + lines * LINE_H + 0.16 for _, _, _, _, label_h, lines in columns)
        if y + row_height > BODY_BOTTOM:
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            y = 0.6

        x = MARGIN
        for item, total, label_w, body_w, label_h, lines in columns:
            body_h = lines * LINE_H
            if not item.is_label and item.label_position != "none":
                _add_text(
                    slide,
                    item.label,
                    x=x,
                    y=y,
                    w=(label_w or total),
                    h=(body_h + label_h if item.label_position == "left" else label_h),
                    size=LABEL_SIZE,
                    bold=True,
                    color=BLACK,
                )
            _add_text(
                slide,
                # 纯栏目标题没有正文，标签就是它的全部文字
                item.value if not item.is_label else item.label,
                x=x + label_w,
                y=y + label_h,
                w=body_w if label_w else total,
                h=body_h,
                size=BODY_SIZE if not item.is_label else LABEL_SIZE,
                bold=item.is_label,
                color=BLACK if item.is_label else DARK,
            )
            x += total + GUTTER
        y += row_height + 0.12


def generate_pptx(
    *,
    title: str,
    fields: dict[str, Any],
    field_labels: dict[str, str],
    content: str,
    layout_fields: list[LayoutField],
    info_line: str = "",
) -> bytes:
    """有栏目就按模板行结构逐行铺页；没有栏目结构时正文单独一页。"""
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    prs.core_properties.title = title
    prs.core_properties.author = "智慧教案"

    _cover(prs, title, info_line)

    rows = _filled_rows(layout_fields, fields, field_labels)
    if rows:
        _append_body(prs, rows)
    elif content:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _add_text(
            slide,
            field_value_to_text(content),
            x=MARGIN,
            y=BODY_TOP,
            w=TEXT_W,
            h=BODY_BOTTOM - BODY_TOP,
            size=INFO_SIZE,
        )

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()
