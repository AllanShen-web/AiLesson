"""教案 → Word（.docx）。

纸面口径对齐教育局下发的模板原文件：标题是表外独立段落（黑体 18pt 居中），表格从第一字段行
直接开始；表宽取模板的 8140 缇并整表水平居中，边框全黑细线；栏目标签用新宋体 11pt 加粗居中，
贯穿格与三字以上的窄标签竖排；正文仿宋 10.5pt 两端对齐、靠单元格顶部排（同一行里内容多的
字段会把行撑高，内容少的若仍垂直居中就会离自己的标签很远）。

python-docx 没有覆盖到这几项 OOXML 能力（文字方向、表级边框、run 底纹），
所以对应处直接写元素，其余都走公开 API。
"""

from __future__ import annotations

import io
import re
from typing import Any

from docx import Document
from docx.document import (
    Document as DocxDocument,  # 前者是工厂函数，后者才是它返回的类：调用用前者，签名用后者
)
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml.parser import OxmlElement
from docx.shared import Pt, RGBColor, Twips
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from ailan.ai_chat.files.layout import (
    GRID_COL_WIDTH,
    GRID_COLUMNS,
    PAGE_HEIGHT,
    PAGE_MARGIN_LR,
    PAGE_MARGIN_TB,
    PAGE_WIDTH,
    LayoutField,
    TableCellSpec,
    TableRowSpec,
    build_segments,
    build_table_cell_specs,
)
from ailan.ai_chat.files.rich_text import (
    InlineRun,
    html_to_paragraphs,
    is_html_content,
    md_inline_to_runs,
    to_text,
)

HEI = "黑体"
LABEL_FONT = "新宋体"
BODY_FONT = "仿宋"
CODE_FONT = "Courier New"
CODE_COLOR = "C7254E"
CODE_FILL = "F9F2F4"

LABEL_SIZE_PT = 11
BODY_SIZE_PT = 10.5
TOP_LABEL_SIZE_PT = 10.5
TITLE_SIZE_PT = 18

# 一行的近似高度（twips）：10.5pt 正文行 + 段间距，用于「高度（行）」的最小行高
ROW_LINE_TWIPS = 320

_HEADING = re.compile(r"^(#{1,4})\s+(.+)$")
_RULE = re.compile(r"^[-*_]{3,}$")
_UL = re.compile(r"^[-*]\s+")
_OL = re.compile(r"^(\d+)\.\s+(.+)$")


# ── python-docx 没暴露的几项格式 ──


def _set_font(run: Any, name: str) -> None:
    """同时设置拉丁与东亚字形名：只设 `font.name` 时中文仍走主题的默认字体。"""
    run.font.name = name
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:eastAsia"), name)


def _set_shading(run: Any, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    run._element.get_or_add_rPr().append(shd)


def _set_vertical_text(cell: _Cell) -> None:
    """竖排文字方向（模板原文件用 tbRlV）：窄列标签与竖排大标题靠它才不会横着挤成一团。"""
    direction = OxmlElement("w:textDirection")
    direction.set(qn("w:val"), "tbRlV")
    cell._tc.get_or_add_tcPr().append(direction)


def _set_table_borders(table: Table) -> None:
    """边框只挂在表级（模板原文件即 tblBorders single 0.5pt 全黑），单元格不再各自描边。"""
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "000000")
        borders.append(element)
    table._tbl.tblPr.append(borders)


def _set_fixed_layout(table: Table) -> None:
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(layout)


def _set_bottom_border(paragraph: Paragraph) -> None:
    """标签在上的标签段落：最后一行标签下描一条与表格线同规格的细线，
    补回原来「标签行 / 填写行」之间的那道分隔线。"""
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:color"), "000000")
    bottom.set(qn("w:space"), "2")
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)


# ── 行内样式片段 → docx run ──


def _add_run(paragraph: Paragraph, run: InlineRun, size_pt: float, font: str | None = None) -> None:
    added = paragraph.add_run(run.text)
    added.font.size = Pt(size_pt)
    added.bold = run.bold or None
    added.italic = run.italics or None
    if run.underline:
        added.underline = True
    if run.strike:
        added.font.strike = True
    if run.code:
        added.font.name = CODE_FONT
        added.font.color.rgb = RGBColor.from_string(CODE_COLOR)
        _set_shading(added, CODE_FILL)
    elif font:
        _set_font(added, font)


def _runs_to_paragraph(container: Any, runs: list[InlineRun], size_pt: float, font: str | None) -> Paragraph:
    paragraph = container.add_paragraph()
    for run in runs:
        _add_run(paragraph, run, size_pt, font)
    return paragraph


# ── markdown / HTML 正文 → docx（助手只给一段正文、没有栏目结构时） ──


def append_markdown(document: DocxDocument, markdown: str) -> None:
    """把 markdown 落成段落与表格：标题、引用、列表、代码块、分隔线各按模板字号排版。"""
    lines = markdown.split("\n")
    index = 0
    while index < len(lines):
        trimmed = lines[index].strip()
        if not trimmed:
            index += 1
            continue

        if trimmed.startswith("```"):
            code: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code.append(lines[index])
                index += 1
            index += 1
            paragraph = document.add_paragraph()
            _indent_left(paragraph)
            run = paragraph.add_run("\n".join(code))
            run.font.name = CODE_FONT
            run.font.size = Pt(10)
            _set_shading(run, "F5F5F5")
            continue

        if trimmed.startswith("|") and trimmed.endswith("|"):
            rows: list[str] = []
            while index < len(lines) and (row := lines[index].strip()).startswith("|") and row.endswith("|"):
                rows.append(row)
                index += 1
            _append_markdown_table(document, rows)
            continue

        heading = _HEADING.match(trimmed)
        if heading:
            level = len(heading.group(1))
            paragraph = document.add_heading("", level=level)
            run = paragraph.add_run(heading.group(2))
            run.bold = True
            run.font.size = Pt({1: 18, 2: 15, 3: 13, 4: 12}[level])
            _set_font(run, HEI)
            continue

        if trimmed.startswith("> "):
            paragraph = document.add_paragraph()
            _indent_left(paragraph)
            run = paragraph.add_run(re.sub(r"^>\s?", "", trimmed))
            run.italic = True
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor.from_string("666666")
            continue

        if _RULE.match(trimmed):
            paragraph = document.add_paragraph()
            run = paragraph.add_run("─" * 40)
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string("CCCCCC")
            continue

        bullet = _UL.match(trimmed)
        if bullet:
            paragraph = document.add_paragraph()
            _indent_hanging(paragraph)
            paragraph.add_run("•  ").font.size = Pt(12)
            _add_inline_markdown(paragraph, bullet.string[2:] if bullet.string else trimmed)
            continue

        ordered = _OL.match(trimmed)
        if ordered:
            paragraph = document.add_paragraph()
            _indent_hanging(paragraph)
            paragraph.add_run(f"{ordered.group(1)}.  ").font.size = Pt(12)
            _add_inline_markdown(paragraph, ordered.group(2))
            continue

        paragraph = document.add_paragraph()
        _add_inline_markdown(paragraph, trimmed)
        index += 1


def _add_inline_markdown(paragraph: Paragraph, text: str) -> None:
    runs = md_inline_to_runs(text)
    if not runs:
        paragraph.add_run(text).font.size = Pt(12)
        return
    for run in runs:
        _add_run(paragraph, run, 12)


def _indent_left(paragraph: Paragraph) -> None:
    paragraph.paragraph_format.left_indent = Twips(432)


def _indent_hanging(paragraph: Paragraph) -> None:
    paragraph.paragraph_format.left_indent = Twips(576)
    paragraph.paragraph_format.first_line_indent = Twips(-288)


def _append_markdown_table(document: DocxDocument, rows: list[str]) -> None:
    """markdown 表格 → docx 表格：分隔行（`| --- | --- |`）只是语法，不占行。"""
    data = [row for row in rows if not re.match(r"^\|[\s\-:|]+\|$", row)]
    if not data:
        return
    cells = [[cell.strip() for cell in row.strip("|").split("|")] for row in data]
    columns = max(len(row) for row in cells)
    table = document.add_table(rows=len(cells), cols=columns)
    table.style = "Table Grid"
    for row_index, row in enumerate(cells):
        for col_index in range(columns):
            paragraph = table.cell(row_index, col_index).paragraphs[0]
            _add_inline_markdown(paragraph, row[col_index] if col_index < len(row) else "")
            if row_index == 0:
                _shade_cell(paragraph, "F0F0F0")


def _shade_cell(paragraph: Paragraph, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), fill)
    paragraph._p.get_or_add_pPr().append(shading)


def append_html(document: DocxDocument, html: str) -> None:
    """HTML 正文 → 段落：块级边界已经在解析阶段还原成段落，这里只负责落字号。"""
    paragraphs = html_to_paragraphs(html)
    if not paragraphs and html.strip():
        paragraphs = [[InlineRun(text=html)]]
    for runs in paragraphs:
        paragraph = document.add_paragraph()
        for run in runs:
            _add_run(paragraph, run, BODY_SIZE_PT + 1.5)


# ── 教案表格 ──


def _label_paragraph(container: Any, runs: list[InlineRun]) -> Paragraph:
    """栏目/左侧标签段落：新宋体 11pt 加粗居中（模板原字号）。"""
    paragraph = _runs_to_paragraph(container, runs, 11, LABEL_FONT)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in paragraph.runs:
        run.bold = True
    return paragraph


def _top_label_paragraph(cell: _Cell, runs: list[InlineRun], with_divider: bool) -> Paragraph:
    """标签在上的标签段落：仿宋 10.5pt 加粗居中、行距固定 15pt（模板原样式）。"""
    paragraph = _runs_to_paragraph(cell, runs, TOP_LABEL_SIZE_PT, BODY_FONT)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in paragraph.runs:
        run.bold = True
    paragraph.paragraph_format.line_spacing = Pt(15)
    if with_divider:
        _set_bottom_border(paragraph)
    return paragraph


def _body_paragraph(cell: _Cell, runs: list[InlineRun]) -> Paragraph:
    """正文段落：仿宋 10.5pt 两端对齐；样式标记已在解析阶段从 HTML/markdown 提取好。"""
    paragraph = _runs_to_paragraph(cell, runs, BODY_SIZE_PT, BODY_FONT)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return paragraph


def _write_cell(cell: _Cell, spec: TableCellSpec) -> None:
    """把一个单元格规格写进 docx 单元格：竖排标签合成单段交给 tbRlV 竖着排。"""
    _clear(cell)

    if spec.is_label:
        if spec.vertical_text:
            label = "".join(run.text for line in spec.lines for run in line)
            _label_paragraph(cell, [InlineRun(text=label)])
            _set_vertical_text(cell)
        else:
            for runs in spec.lines:
                _label_paragraph(cell, runs)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        return

    label_lines = spec.label_lines or []
    for index, runs in enumerate(label_lines):
        _top_label_paragraph(cell, runs, index == len(label_lines) - 1)
    for runs in spec.lines:
        _body_paragraph(cell, runs)
    # 正文格靠顶：同一行里内容多的字段会把行撑高，内容少的若仍按垂直居中就会离自己的标签很远
    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP


def _clear(cell: _Cell) -> None:
    """清掉建表时自带的空段落，避免每个格子多出一行行高。"""
    for paragraph in list(cell.paragraphs):
        paragraph._p.getparent().remove(paragraph._p)


def _append_table(document: DocxDocument, specs: list[TableRowSpec]) -> None:
    """一段单元格规格 → 一张表格（分隔符把字段切成的每一段各一张）。

    合并格子的顺序是先横向并栏、再纵向贯穿，所以先把 24 栏网格的 `_Cell` 对象全部拿到手：
    合并会改变每行的单元格列表，之后再按行列号取就对不上网格了。
    """
    table = document.add_table(rows=len(specs), cols=GRID_COLUMNS)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_fixed_layout(table)
    _set_table_borders(table)
    for column in table.columns:
        column.width = Twips(GRID_COL_WIDTH)

    grid = [
        [table.cell(row_index, col_index) for col_index in range(GRID_COLUMNS)]
        for row_index in range(len(specs))
    ]

    for row_index, spec in enumerate(specs):
        cursor = spec.covered_span
        row_cells: list[tuple[_Cell, TableCellSpec]] = []
        for cell_spec in spec.cells:
            left = grid[row_index][min(cursor, GRID_COLUMNS - 1)]
            right_col = min(cursor + cell_spec.span - 1, GRID_COLUMNS - 1)
            bottom_row = min(row_index + (cell_spec.row_span or 1) - 1, len(specs) - 1)
            target = grid[bottom_row][right_col] if bottom_row != row_index or right_col != cursor else left
            merged = left if target is left else left.merge(target)
            row_cells.append((merged, cell_spec))
            cursor += cell_spec.span

        # 贯穿格的高度摊在它罩住的行上，不参与本行的最小行高
        own = [cell_spec for _, cell_spec in row_cells if not (cell_spec.row_span or 1) > 1]
        # 本行需要「留白」到几行：只按模板设的「高度（行）」算，正文本身比它高时不再加码。
        # 早先取的是 max(正文行数, 设定行数)，正文几十行时等于向 Word 声明一个比一页还高的最小行高，
        # Word 会把整行挪到下一页去满足这个高度，当前页就留下大片空白。
        pad_lines = max(
            (
                len(cell_spec.label_lines or []) + cell_spec.min_lines
                for cell_spec in own
                if cell_spec.min_lines
            ),
            default=0,
        )
        row = table.rows[row_index]
        if pad_lines > 0:
            row.height = Twips(pad_lines * ROW_LINE_TWIPS)
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

    # 内容要等所有合并落定后再写：合并会把多个格子的段落串到一起
    for row_index, spec in enumerate(specs):
        cursor = spec.covered_span
        for cell_spec in spec.cells:
            cell = table.cell(row_index, min(cursor, GRID_COLUMNS - 1))
            if cell_spec.is_label or not cell_spec.row_span or cell_spec.row_span <= 1:
                _write_cell(cell, cell_spec)
            cursor += cell_spec.span


def _setup_page(document: DocxDocument) -> None:
    """显式 A4 与 1.27cm 页边距：与表格宽度假设一致，Word 与预览端才按同一张纸排版。"""
    section = document.sections[0]
    section.page_width = Twips(PAGE_WIDTH)
    section.page_height = Twips(PAGE_HEIGHT)
    section.top_margin = Twips(PAGE_MARGIN_TB)
    section.bottom_margin = Twips(PAGE_MARGIN_TB)
    section.left_margin = Twips(PAGE_MARGIN_LR)
    section.right_margin = Twips(PAGE_MARGIN_LR)


def generate_docx(
    *,
    title: str,
    fields: dict[str, Any],
    field_labels: dict[str, str],
    content: str,
    layout_fields: list[LayoutField],
) -> bytes:
    """有栏目内容时按模板表格排版；否则按 HTML/markdown 正文输出（助手直接写的一段话）。"""
    document = Document()
    _setup_page(document)

    if fields:
        build_lesson_plan(document, title, layout_fields, fields, field_labels)
    elif is_html_content(content):
        append_html(document, content)
    elif content:
        append_markdown(document, content)
    else:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(to_text(content))
        run.font.size = Pt(12)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_lesson_plan(
    document: DocxDocument,
    title: str,
    fields: list[LayoutField],
    values: dict[str, Any],
    field_labels: dict[str, str],
) -> None:
    """标题段落 + 模板表格，依次追加进给定 document。"""
    tables = [
        (build_table_cell_specs(title, segment.rows)[1:], segment.page_break)
        for segment in build_segments(fields, values, field_labels)
    ]

    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_paragraph.add_run(title or "教案")
    run.font.size = Pt(TITLE_SIZE_PT)
    _set_font(run, HEI)

    for specs, page_break in tables:
        if not specs:
            continue
        if len(document.element.body) > 2:
            # 相邻的两张 Table 会被 Word 合并成一张，中间必须隔一个段落；
            # 连续分隔符那条边界用分页段落——它既是隔开表格的段落，又真换页
            spacer = document.add_paragraph()
            if page_break:
                spacer.add_run().add_break(WD_BREAK.PAGE)
        _append_table(document, specs)
